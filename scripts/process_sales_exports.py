"""
Consolidates Kauai County Real Property Tax sales exports (manually pulled
by the user from the county's qPublic/Schneider Corp CAMA system -- no
self-service bulk export exists there, confirmed 2026-10-15) into one
table joined to this project's own parcel/zoning/tract pipeline, and
computes real observed $/building-sqft by zoning district and by tract.

This is meant to validate, and eventually replace, the single Lihu'e
Town Core reference sale price + ACS-median-home-value tract scaling
used in kauai_development_fee_reduction_feasibility() -- real 2026
transactions are a stronger revenue input than a 2023 secondary-sourced
pro forma scaled by a Census proxy.

PARCEL ID JOIN KEY, reverse-engineered and confirmed 2026-10-15: these
exports' 12-digit "Parcel ID" is Section(2) + Plat(3) + Parcel(3) +
CPR/unit(4), with NO leading zone digit. build_model.py's parcel_level
tmk_txt is Zone(1) + Section(2) + Plat(3) + Parcel(3), 9 digits, zone
always "4" for Kauai. So tmk_txt[1:] (8 digits) == Parcel ID[:8] (8
digits) is the real join key -- confirmed exact match on all sales
records against data/processed/kauai_parcel_level.csv, now spanning
Feb 2025-Jun 2026 (17 months, pulled to match the Zillow ZORI rent
series' own history on the Existing Conditions "over time" chart).
This drops the CPR/condo-unit suffix, same granularity loss as
parcel_level's own "one row per physical parcel" structure.

Run after build_county("kauai") (needs kauai_parcel_level.csv) and after
placing monthly export files in data/raw/sales_exports/ (gitignored --
not self-service, so there's no fetch_data.py layer for this; the user
supplies new months as they're able to pull them).
"""
import glob
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).parent))
import build_model as bm

SALES_DIR = Path(__file__).parent.parent / "data" / "raw" / "sales_exports"
OUT_PATH = Path(__file__).parent.parent / "data" / "processed" / "kauai_sales_2026_joined.csv"


def _read_export(path: Path) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header, data = rows[0], rows[1:]
    df = pd.DataFrame(data, columns=[h.strip() for h in header])
    return df


def load_kauai_sales() -> pd.DataFrame:
    files = sorted(glob.glob(str(SALES_DIR / "kauai_sales_20*.xlsx")))
    if not files:
        raise RuntimeError(
            f"No Kauai sales exports found in {SALES_DIR} -- these are manually pulled from "
            "the county's qPublic site (no bulk self-service export exists), so there's "
            "nothing to fetch automatically. Ask for a fresh monthly export and place it here."
        )
    frames = [_read_export(Path(f)) for f in files]
    sales = pd.concat(frames, ignore_index=True)

    sales["key8"] = sales["Parcel ID"].astype(str).str[:8]
    sales["sale_price_usd"] = (
        sales["Sale Price"].astype(str).str.replace(r"[$,]", "", regex=True).astype(float)
    )
    sales["sale_date"] = pd.to_datetime(sales["Sale Date"])
    sales["building_sqft"] = pd.to_numeric(sales["Square Ft"], errors="coerce")
    sales["price_per_sqft"] = sales["sale_price_usd"] / sales["building_sqft"]
    return sales


def join_to_parcels(sales: pd.DataFrame) -> pd.DataFrame:
    parcel_path = bm.OUT_DIR / "kauai_parcel_level.csv"
    if not parcel_path.exists():
        raise FileNotFoundError(f"Run build_county('kauai') first -- {parcel_path} not found")
    parcels = pd.read_csv(parcel_path, dtype={"tmk_txt": str})
    parcels["key8"] = parcels["tmk_txt"].str[1:]

    joined = sales.merge(
        parcels[["key8", "GEOID", "zoning", "taxclass"]], on="key8", how="left"
    )
    unmatched = joined["GEOID"].isna().sum()
    if unmatched:
        print(f"  warning: {unmatched}/{len(joined)} sales records did not match a parcel "
              "(likely a CPR/condo unit on a multi-unit TMK, or a parcel outside this layer's coverage)")
    return joined


def summarize(joined: pd.DataFrame) -> None:
    valid = joined[
        (joined["building_sqft"] > 0)
        & joined["price_per_sqft"].notna()
        & (joined["taxclass"].isin(bm.COUNTY_CONFIG["kauai"]["existing_residential_taxclasses"]))
    ]
    print(f"\n{len(joined)} total sales records, {len(valid)} usable residential land+building sales "
          f"with building sqft (Feb 2025-Jun 2026)")
    print(f"Median $/building sqft, all usable sales: ${valid['price_per_sqft'].median():,.0f}")
    print("\nBy zoning district (n >= 5):")
    by_zone = valid.groupby("zoning")["price_per_sqft"].agg(["median", "count"])
    print(by_zone[by_zone["count"] >= 5].sort_values("count", ascending=False))
    print("\nBy tract (n >= 5) -- this is the granularity the feasibility pro forma would need:")
    by_tract = valid.groupby("GEOID")["price_per_sqft"].agg(["median", "count"])
    qualifying = by_tract[by_tract["count"] >= 5]
    print(qualifying if len(qualifying) else "  none -- no tract has 5+ usable sales over the full window")
    print(f"\n{len(by_tract)} of 17 tracts have ANY usable sale; "
          f"max per tract: {by_tract['count'].max() if len(by_tract) else 0}")


def main() -> None:
    sales = load_kauai_sales()
    joined = join_to_parcels(sales)
    joined.to_csv(OUT_PATH, index=False)
    print(f"{len(joined)} records -> {OUT_PATH}")
    summarize(joined)


if __name__ == "__main__":
    main()
